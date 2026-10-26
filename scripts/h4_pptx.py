# Stage H.4: the presentation as an editable PowerPoint deck

# Imports
import json
import sys
from pathlib import Path

import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from f1radio.config import load_config, path

cfg = load_config()
R, FIG = path(cfg, "results"), path(cfg, "figures")
OUT = FIG.parent / "F1_radio_presentation.pptx"


def J(rel):
    return json.loads((R / rel).read_text())


attr, wer, txt = J("stage_a/attribution_eval.json"), J("stage_b/wer.json"), J("stage_e/text_metrics.json")
aro, ft, abl = J("stage_e/arousal_metrics.json"), J("stage_e/wavlm_ft_metrics.json"), J("stage_e/wavlm_ablation.json")
t2r, t2n, t4 = J("stage_e/t2_rules.json"), J("stage_e/t2_neural.json"), J("stage_g/t4_results.json")
llm, lc = J("stage_e/llm_zeroshot.json"), J("stage_e/learning_curve.json")
resid = J("stage_d/residual_model_diagnostics.json")["summary"]
lag, dedup, join = J("stage_d/timestamp_offset.json"), J("stage_a/dedup_report.json"), J("stage_d/join_summary.json")
clips = pd.read_parquet(path(cfg, "interim") / "clips.parquet")
fetch = pd.read_csv(R / "stage_d" / "fetch_log.csv")
cv = txt["protocols"]["cv5"]

# Design system
DARK, RED, INK, MUTED = RGBColor(0x16, 0x18, 0x1D), RGBColor(0xC8, 0x10, 0x2E), RGBColor(0x16, 0x18, 0x1D), RGBColor(0x5B, 0x60, 0x70)
TINT, LINE, WHITE = RGBColor(0xF3, 0xF4, 0xF6), RGBColor(0xE3, 0xE5, 0xEA), RGBColor(0xFF, 0xFF, 0xFF)
GOOD, BAD, AMBER = RGBColor(0x1D, 0x7A, 0x46), RGBColor(0xB4, 0x23, 0x2C), RGBColor(0x9A, 0x6B, 0x00)
BLUE, ORANGE, AQUA, GREY = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0xA9, 0xAD, 0xB8)
HEAD, BODY = "Cambria", "Calibri"
W, H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
state = {"section": "", "n": 0}


def new_slide(dark=False, section=None):
    s = prs.slides.add_slide(BLANK)
    state["n"] += 1
    if section is not None:
        state["section"] = section
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK if dark else WHITE
    if state["n"] > 1:
        tb = s.shapes.add_textbox(Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.3))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = str(state["n"])
        r.font.size, r.font.name = Pt(10), BODY
        r.font.color.rgb = GREY if dark else MUTED
        if state["section"] and not dark:
            tb = s.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(6), Inches(0.3))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = state["section"].upper()
            r.font.size, r.font.name, r.font.bold = Pt(9), BODY, True
            r.font.color.rgb = MUTED
    return s


def text(s, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", italic)
            r.font.color.rgb = o.get("color", color)
    return tb


def title(s, t, sub=None, dark=False):
    text(s, Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.8), t, size=32, font=HEAD, bold=True,
         color=WHITE if dark else INK)
    if sub:
        text(s, Inches(0.6), Inches(1.15), Inches(12.1), Inches(0.5), sub, size=16, color=GREY if dark else MUTED)


def bullets(s, x, y, w, h, items, size=17, color=INK, gap=8):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.28))))
        pPr.set("indent", str(-Emu(Inches(0.28))))
        bu = etree.SubElement(pPr, qn("a:buFont"))
        bu.set("typeface", "Arial")
        bc = etree.SubElement(pPr, qn("a:buChar"))
        bc.set("char", "•")
        p.space_after = Pt(gap)
        for t, o in ([(it, {})] if isinstance(it, str) else it):
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size = BODY, Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(s, x, y, w, h, fill=TINT, line=None, radius=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def card(s, x, y, w, h, head, body, fill=TINT, head_color=INK, size=14):
    box(s, x, y, w, h, fill)
    text(s, x + Inches(0.18), y + Inches(0.14), w - Inches(0.36), Inches(0.45), head, size=17, bold=True, color=head_color)
    text(s, x + Inches(0.18), y + Inches(0.6), w - Inches(0.36), h - Inches(0.7), body, size=size, color=MUTED)


def stat(s, x, y, w, h, value, label, color=RED, fill=TINT):
    box(s, x, y, w, h, fill)
    text(s, x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), Inches(0.85), value, size=36, font=HEAD, bold=True, color=color)
    text(s, x + Inches(0.2), y + Inches(0.95), w - Inches(0.4), h - Inches(1.0), label, size=13, color=MUTED)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def image(s, name, x, y, w=None, h=None):
    p = FIG / name
    return s.shapes.add_picture(str(p), x, y, width=w, height=h)


def bar_chart(s, x, y, w, h, cats, vals, colors, fmt="0.00", horizontal=True, vmax=None, label_size=12, title_txt=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("v", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = s.shapes.add_chart(kind, x, y, w, h, cd).chart
    ch.has_legend = False
    ch.font.name, ch.font.size = BODY, Pt(12)
    ch.font.color.rgb = INK
    if title_txt:
        ch.has_title = True
        ch.chart_title.text_frame.text = title_txt
        tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
        tr.font.size, tr.font.bold, tr.font.name = Pt(13), True, BODY
        tr.font.color.rgb = MUTED
    else:
        ch.has_title = False
    pl = ch.plots[0]
    pl.gap_width = 55
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(label_size), True
    for i, pt in enumerate(pl.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i % len(colors)]
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    va.minimum_scale = 0
    if vmax:
        va.maximum_scale = vmax
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(12)
    ca.format.line.color.rgb = LINE
    if horizontal:
        ca.reverse_order = True
    return ch


def clustered(s, x, y, w, h, cats, series, colors, fmt="0%", vmax=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, cd).chart
    ch.font.name, ch.font.size = BODY, Pt(12)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    pl = ch.plots[0]
    pl.gap_width, pl.overlap = 70, -10
    pl.has_data_labels = True
    pl.data_labels.number_format, pl.data_labels.number_format_is_linked = fmt, False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(10)
    for ser, col in zip(pl.series, colors):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    va.minimum_scale = 0
    if vmax:
        va.maximum_scale = vmax
    ch.category_axis.tick_labels.font.size = Pt(12)
    ch.category_axis.format.line.color.rgb = LINE
    return ch


def table(s, x, y, w, rows, col_w, header_fill=DARK, size=13, row_h=0.42, fills=None):
    t = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.name, r.font.size = BODY, Pt(size)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = header_fill
                r.font.bold, r.font.color.rgb = True, WHITE
            else:
                c.fill.fore_color.rgb = (fills[i][j] if fills and fills[i][j] is not None else (WHITE if i % 2 else TINT))
                r.font.color.rgb = INK
    return t


def arrow(s, x1, y1, x2, y2, color=MUTED):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def section(num, name, blurb):
    s = new_slide(dark=True, section=name)
    text(s, Inches(0.8), Inches(2.3), Inches(3), Inches(1.4), f"{num:02d}", size=80, font=HEAD, bold=True, color=RED)
    text(s, Inches(0.8), Inches(3.6), Inches(11), Inches(1), name, size=40, font=HEAD, bold=True, color=WHITE)
    text(s, Inches(0.8), Inches(4.55), Inches(10.5), Inches(1), blurb, size=18, color=GREY)
    return s


def pct(x):
    return f"{100 * x:.0f}%"


# 
# Title
s = new_slide(dark=True)
text(s, Inches(0.8), Inches(1.6), Inches(11), Inches(0.5), "NLP + DEEP LEARNING PROJECT", size=14, bold=True, color=RED)
text(s, Inches(0.8), Inches(2.1), Inches(11.5), Inches(2.2), ["What drivers claim,", "and what the car says"], size=54,
     font=HEAD, bold=True, color=WHITE)
text(s, Inches(0.8), Inches(4.35), Inches(11), Inches(0.8),
     "Understanding Formula 1 team radio with speech and language models, and checking it against the car's telemetry",
     size=20, color=GREY)
text(s, Inches(0.8), Inches(5.8), Inches(11), Inches(0.5), "Presented by: ____________________", size=16, color=GREY)
notes(s, "Open with the one-line idea: we listen to F1 radio, work out what drivers claim, and check whether the car's data backs them up.")

# Agenda
s = new_slide(section="")
title(s, "Agenda")
agenda = ["Problem statement", "The dataset", "Literature review", "Methodology", "Problems faced, and how we solved them",
          "Results", "Limitations and future scope", "Conclusion"]
for i, a in enumerate(agenda):
    col, row = i // 4, i % 4
    x, y = Inches(0.8 + col * 6.2), Inches(1.7 + row * 1.25)
    box(s, x, y, Inches(5.7), Inches(1.0))
    text(s, x + Inches(0.25), y + Inches(0.12), Inches(0.9), Inches(0.8), f"{i + 1:02d}", size=30, font=HEAD, bold=True,
         color=RED, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(1.2), y + Inches(0.12), Inches(4.3), Inches(0.8), a, size=20, anchor=MSO_ANCHOR.MIDDLE)

# 01 Problem
section(1, "Problem statement", "What are we trying to understand, and why is it hard?")

s = new_slide(section="Problem statement")
title(s, "The problem", "F1 drivers talk to their engineers over the radio at 300 km/h")
bullets(s, Inches(0.6), Inches(1.9), Inches(6.6), Inches(4.5), [
    "During a race, drivers keep telling their engineers what's wrong: “no grip”, “he pushed me off”, “I've lost power”.",
    "Teams and TV commentators take these at face value, but nobody really checks whether the car's own data agrees.",
    [("Our goal ", {"bold": True}), ("was a system that listens to one radio message and answers four questions about it on its own.", {})],
    [("Why it matters: ", {"bold": True}), ("it's a rare case where spoken language can be checked against an independent, physical record.", {})],
], size=18)
box(s, Inches(7.6), Inches(1.9), Inches(5.1), Inches(4.6), DARK)
text(s, Inches(7.9), Inches(2.1), Inches(4.6), Inches(0.5), "ONE RADIO MESSAGE", size=12, bold=True, color=RED)
text(s, Inches(7.9), Inches(2.55), Inches(4.6), Inches(0.9), "“No grip at the rear!”", size=26, font=HEAD, bold=True, color=WHITE, italic=True)
for i, (q, a) in enumerate([("Who said it?", "the driver, not the engineer"), ("What is claimed?", "tyre grip / degradation"),
                            ("How worked up?", "raised voice"), ("Does the car agree?", "were lap times dropping?")]):
    y = Inches(3.6 + i * 0.7)
    text(s, Inches(7.9), y, Inches(2.2), Inches(0.6), q, size=15, bold=True, color=WHITE)
    text(s, Inches(10.1), y, Inches(2.5), Inches(0.6), a, size=15, color=GREY)
notes(s, "Use the dark card as the running example: one message, four questions.")

s = new_slide(section="Problem statement")
title(s, "Four research questions, and why each is hard")
qs = [("T1  Claim type", "What is the driver complaining about: tyres, traffic, pace, an incident, the engine?",
       "The messages are short, noisy and full of jargon, and there were almost no labels"),
      ("T2  Telemetry check", "Does the car's own sensor data back up what the driver said?",
       "Public telemetry misses many sensors, and TV picks the dramatic moments"),
      ("T3  Arousal", "How worked up does the driver sound, from the words and from the voice?",
       "Engine noise, radio compression and g-forces all change how a voice sounds"),
      ("T4  Consequences", "Does a worked-up driver then go faster or slower?",
       "Lap times swing by about 0.7 s anyway, so any effect may be tiny")]
for i, (h, q, why) in enumerate(qs):
    x = Inches(0.6 + i * 3.08)
    box(s, x, Inches(1.7), Inches(2.85), Inches(4.5))
    text(s, x + Inches(0.2), Inches(1.85), Inches(2.5), Inches(0.5), h, size=19, bold=True, color=RED)
    text(s, x + Inches(0.2), Inches(2.45), Inches(2.5), Inches(1.9), q, size=16)
    text(s, x + Inches(0.2), Inches(4.45), Inches(2.5), Inches(0.4), "Why it's hard", size=13, bold=True, color=MUTED)
    text(s, x + Inches(0.2), Inches(4.85), Inches(2.5), Inches(1.6), why, size=14, color=MUTED)
text(s, Inches(0.6), Inches(6.35), Inches(12), Inches(0.4),
     "And underneath all four, every clip mixes the driver's and the engineer's voices, with nothing to say who's talking.", size=15, italic=True, color=MUTED)

# 02 Dataset
section(2, "The dataset", "Real broadcast radio, joined to real car telemetry")

s = new_slide(section="The dataset")
title(s, "Radio data: F1 Team Radio (Hugging Face)", "Broadcast clips with audio and transcripts, 2018–2025")
n_raw, n_keep = dedup["n_input"], dedup["n_kept"]
hours = clips["duration_s"].sum() / 3600
for i, (v, l) in enumerate([(f"{n_raw:,}", "radio clips in the dataset"), (f"{hours:.0f} h", "of audio"),
                            (str(clips["driver_id"].nunique()), "drivers"), (str(clips["race_id"].nunique()), "Grands Prix")]):
    stat(s, Inches(0.6 + (i % 2) * 3.2), Inches(1.75 + (i // 2) * 1.5), Inches(3.05), Inches(1.38), v, l)
bullets(s, Inches(0.6), Inches(4.85), Inches(6.3), Inches(2.1), [
    "Every clip comes with its audio, a machine transcript, the driver, the race and the time it aired.",
    f"A typical clip lasts about {clips['duration_s'].median():.0f} seconds, and many contain both driver and engineer.",
    "2022 is missing completely and 2025 is thin, so our held-out test season is small.",
    "It's listed as CC BY 4.0, but it's broadcast audio, so we used it for coursework only.",
], size=14, gap=4)
yr = clips.groupby("year").size()
yrs = list(range(2018, 2026))
ch = bar_chart(s, Inches(7.1), Inches(1.7), Inches(5.7), Inches(4.9), [str(y) if y != 2022 else "2022*" for y in yrs],
               [int(yr.get(y, 0)) for y in yrs], [RED if y != 2025 else ORANGE for y in yrs], fmt="#,##0",
               horizontal=False, title_txt="Clips per season (after cleaning)")
text(s, Inches(7.1), Inches(6.6), Inches(5.7), Inches(0.4), "*2022 missing from the source; orange = 2025, the test season",
     size=12, color=MUTED, align=PP_ALIGN.CENTER)

s = new_slide(section="The dataset")
title(s, "Inside one clip", "What the raw data looks like, and why it's messy")
box(s, Inches(0.6), Inches(1.8), Inches(6.4), Inches(4.8), TINT)
text(s, Inches(0.85), Inches(1.95), Inches(6), Inches(0.4), "A real transcript, exactly as stored (2020, Albon's channel)", size=13, bold=True, color=MUTED)
text(s, Inches(0.85), Inches(2.45), Inches(5.9), Inches(2.6),
     "“The drink system is leaking a lot. Copy. We'll have a look on the grid. Copy. Yeah, I'm getting soaked here. Yeah, copy. Just make sure we have a spare balaclava.”",
     size=20, italic=True)
box(s, Inches(0.85), Inches(5.1), Inches(5.9), Inches(1.3), WHITE, line=LINE)
text(s, Inches(1.05), Inches(5.2), Inches(5.5), Inches(1.1), [[("Who said which sentence? ", {"bold": True, "color": RED}),
     ("The file doesn't say. Our pipeline has to work it out from the voices.", {})]], size=16)
table(s, Inches(7.4), Inches(1.8), Inches(5.3), [
    ["Field", "Meaning"], ["driver_id", "whose radio channel (not who speaks)"], ["race_id", "which Grand Prix"],
    ["message_timestamp", "UTC time the clip aired"], ["audio", "MP3, 44–48 kHz"], ["transcription", "machine ASR (with errors)"]],
    [2.0, 3.3], size=13)
bullets(s, Inches(7.4), Inches(4.75), Inches(5.3), Inches(1.8), [
    "Two people share one recording, and nothing marks where one stops and the other starts.",
    "The machine transcript mangles jargon: “SuperSalt” instead of supersoft.",
], size=15)

s = new_slide(section="The dataset")
title(s, "Telemetry: what the car tells us", "Fetched with FastF1 and joined to each radio clip by car number and time")
ok_races = int(fetch["status"].isin(["ok", "ok_cached", "ok_no_car_data"]).sum())
card(s, Inches(0.6), Inches(1.8), Inches(3.9), Inches(3.3), "Available (public)",
     "Speed, throttle, brake on/off, RPM, gear, DRS (~3.7 Hz)\nLap and sector times, gaps, positions\nTyre compound and age, pit stops\nRace-control messages, weather")
card(s, Inches(4.75), Inches(1.8), Inches(3.9), Inches(3.3), "Not available",
     "Brake pressure or temperature\nTyre temperatures\nSteering angle, ride height\nFuel load\n→ claims about these can't be checked")
for i, (v, l) in enumerate([(f"{ok_races}", "of 149 races with telemetry"), (f"{join['race_phase_by_year'] and sum(join['race_phase_by_year'].values()):,}", "radio clips matched to a racing lap"),
                            (f"{resid['clean_laps']:,}", "clean laps for the lap-time model")]):
    stat(s, Inches(8.95), Inches(1.8 + i * 1.6), Inches(3.75), Inches(1.45), v, l)
bullets(s, Inches(0.6), Inches(5.35), Inches(8.0), Inches(1.6), [
    "We also built a lap-time model that predicts how fast each lap should be from fuel, tyre age, tyre type and traffic, so we can tell when a driver was genuinely slow.",
    f"On laps it hadn't seen, it's typically off by {resid['oos_rmse_s_median_race']:.2f} s, compared with {resid['naive_driver_mean_rmse_s_median_race']:.2f} s for a simple per-driver average.",
], size=15)

s = new_slide(section="The dataset")
title(s, "Our own labelled data", "Hand-labelled by us, because no labels for these tasks existed")
for i, (v, l) in enumerate([(str(attr["n_labelled_main"] + attr["test_retest"]["n"]), "turns labelled driver vs engineer"),
                            (str(wer["n_clips"]), "clips hand-transcribed (for WER)"),
                            ("450", "utterances labelled: claim type + arousal")]):
    stat(s, Inches(0.6), Inches(1.8 + i * 1.6), Inches(3.6), Inches(1.45), v, l)
cc = txt["claim_counts"]
order = sorted(cc, key=lambda k: -cc[k])
nice = {"NO_CLAIM": "No claim", "GRIP_DEG": "Tyres / grip", "NO_SENSOR": "Can't be checked", "REL_PACE": "Pace vs others",
        "INCIDENT": "Incident", "TRAFFIC": "Traffic", "POWER": "Power / engine"}
bar_chart(s, Inches(4.5), Inches(1.7), Inches(4.4), Inches(4.6), [nice[k] for k in order], [cc[k] for k in order],
          [GREY if cc[k] < 30 else BLUE for k in order], fmt="0", title_txt="Claim types (driver utterances)")
bullets(s, Inches(9.2), Inches(1.9), Inches(3.6), Inches(4.6), [
    "We defined the categories by what a car sensor could actually check.",
    f"{txt['n_not_driver_flagged']} of the 450 turned out to be the engineer, which shows how much the speaker step matters.",
    "Grey classes have fewer than 30 examples, too few to learn from, so we left them out.",
    f"When we re-labelled repeated items, we agreed with ourselves {pct(attr['test_retest']['agreement'])} of the time.",
], size=15)

# 03 Literature
section(3, "Literature review", "What already exists, what we borrowed, and the gap we fill")

s = new_slide(section="Literature review")
title(s, "Related work: speech and language")
table(s, Inches(0.6), Inches(1.5), Inches(12.1), [
    ["Area", "Key work", "What it gives us", "Gap for our problem"],
    ["Speech recognition", "Whisper (Radford et al., 2023)", "Transcribes noisy speech well, and can be nudged with a prompt", "Never tested on compressed race radio full of jargon"],
    ["Speaker diarization", "pyannote (Bredin, 2023); WeSpeaker (Wang et al., 2023)", "Works out who spoke when, and fingerprints each voice", "Tells voices apart, but not which one is the driver"],
    ["Text encoders", "DeBERTa-v3 (He et al., 2023)", "A strong language model we can fine-tune", "Usually needs thousands of labels; we had about 300"],
    ["Large language models", "GPT-3 few-shot (Brown et al., 2020); Qwen2.5 (2024)", "Can classify from written instructions, without training", "Unclear whether that beats fine-tuning on tiny, specialised data"],
    ["Speech representations", "WavLM (Chen et al., 2022); SUPERB layer weighting (Yang et al., 2021)", "Rich audio features, and a way to choose which layers to use", "Not yet tried for arousal in noisy radio"],
], [2.1, 3.3, 3.2, 3.5], size=13, row_h=0.78)

s = new_slide(section="Literature review")
title(s, "Related work: verification, fusion and method")
table(s, Inches(0.6), Inches(1.5), Inches(7.6), [
    ["Area", "Key work", "Our take"],
    ["Claim verification", "FEVER (Thorne et al., 2018)", "Our evidence is sensor data, not text"],
    ["Multimodal fusion", "Constrained fusion (Lee et al., 2026)", "We test gated fusion against simple joining"],
    ["Emotion / arousal", "Circumplex of affect (Russell, 1980)", "Why we say “arousal”, not “stress”"],
    ["Research practice", "Pre-registration (Nosek et al., 2018); p-hacking (Simmons et al., 2011)", "We fixed our decisions before seeing results"],
    ["F1 data", "FastF1; OpenF1", "Where our car data comes from"],
], [2.0, 3.2, 2.4], size=13, row_h=0.72)
box(s, Inches(8.6), Inches(1.5), Inches(4.1), Inches(4.85), DARK)
text(s, Inches(8.85), Inches(1.7), Inches(3.7), Inches(0.5), "THE GAP WE FILL", size=13, bold=True, color=RED)
text(s, Inches(8.85), Inches(2.2), Inches(3.7), Inches(4), [
    "Most fact-checking research checks written claims against written evidence.",
    "",
    "We check spoken claims against a physical record, the car's telemetry, in an area with almost no labels.",
    "",
    "And we test whether any apparent agreement survives a placebo control.",
], size=16, color=WHITE)

# 04 Method
section(4, "Methodology", "How we tried to solve the problem: a six-stage pipeline")

s = new_slide(section="Methodology")
title(s, "The pipeline at a glance")
stages = [("Radio audio", "14,370 clips"), ("Who speaks?", "pyannote + voice clustering"), ("Speech → text", "Whisper large-v3"),
          ("What is claimed?", "DeBERTa · TF-IDF · LLM"), ("How aroused?", "WavLM fine-tuned"), ("Car check", "telemetry rules + neural")]
bw, gap = 1.85, 0.2
for i, (a, b) in enumerate(stages):
    x = Inches(0.6 + i * (bw + gap))
    box(s, x, Inches(2.1), Inches(bw), Inches(2.0), DARK if i in (0,) else TINT)
    text(s, x + Inches(0.12), Inches(2.2), Inches(0.6), Inches(0.5), str(i + 1), size=22, font=HEAD, bold=True, color=RED)
    text(s, x + Inches(0.12), Inches(2.75), Inches(bw - 0.24), Inches(0.7), a, size=16, bold=True, color=WHITE if i == 0 else INK)
    text(s, x + Inches(0.12), Inches(3.4), Inches(bw - 0.24), Inches(0.7), b, size=12, color=GREY if i == 0 else MUTED)
    if i < len(stages) - 1:
        arrow(s, x + Inches(bw) + Inches(0.02), Inches(3.1), x + Inches(bw + gap) - Inches(0.02), Inches(3.1))
bullets(s, Inches(0.6), Inches(4.6), Inches(12), Inches(2.2), [
    "Before running each stage, we wrote down what would count as success, for example speaker accuracy of at least 85%.",
    "Every model is compared with simple baselines: guessing the most common class, a model that never reads the words, and pure chance.",
    "We always keep whole races together, so the same race never shows up in both training and testing.",
], size=17)

s = new_slide(section="Methodology")
title(s, "Speech stages: who spoke, and what they said")
card(s, Inches(0.6), Inches(1.55), Inches(3.95), Inches(2.75), "1. Diarization",
     "pyannote splits each clip into speaker turns, giving 29,134 turns across the 14,370 clips.", size=15)
card(s, Inches(4.7), Inches(1.55), Inches(3.95), Inches(2.75), "2. Driver or engineer?",
     "We group similar voices on each driver's channel, then name each group from what it says: engineers say “box”, drivers say “my tyres”.", size=15)
card(s, Inches(8.8), Inches(1.55), Inches(3.95), Inches(2.75), "3. Speech to text",
     "Whisper large-v3 writes down the words with their timings. We nudged it with a prompt of driver names and F1 jargon, and measured the errors.", size=15)
T = pd.read_parquet(path(cfg, "processed") / "turns_attributed.parquet")
ex = T[(T["clip_id"] == "2020_Austrian_Grand_Prix_ALEALB01_23_20200705_134508") & T["attributable"]].sort_values("start")
ex_dur = float(clips.set_index("id").loc["2020_Austrian_Grand_Prix_ALEALB01_23_20200705_134508", "duration_s"])
text(s, Inches(0.6), Inches(4.5), Inches(12.1), Inches(0.4),
     "Our pipeline's actual output for the clip on slide 8", size=14, bold=True, color=MUTED)
sc = 12.1 / ex_dur
for r in ex.itertuples():
    x0, x1 = Inches(0.6 + r.start * sc), Inches(0.6 + r.end * sc)
    who = "Driver" if r.ens_silk == "DRIVER" else "Engineer" if r.ens_silk == "ENGINEER" else "?"
    box(s, x0, Inches(4.9), x1 - x0, Inches(0.5), RED if who == "Driver" else GREY)
    text(s, x0 + Inches(0.08), Inches(4.92), x1 - x0 - Inches(0.16), Inches(0.46), who, size=13, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0, Inches(5.45), x1 - x0, Inches(0.7), "“" + r.text.strip() + "”", size=12, color=MUTED, italic=True)
text(s, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.5), [[("Every turn was called “driver”, ", {"bold": True, "color": RED}),
     ("but “spare balaclava” is clearly the engineer, and the first turn merges both voices. This is why speaker accuracy is only 67%.", {})]], size=14)

s = new_slide(section="Methodology")
title(s, "Language and voice models")
box(s, Inches(0.6), Inches(1.6), Inches(6.1), Inches(5.0))
text(s, Inches(0.85), Inches(1.75), Inches(5.6), Inches(0.5), "Claim type: three ways to use a language model", size=17, bold=True)
for i, (h, b, col) in enumerate([("TF-IDF + logistic regression", "counts the words: our simple baseline", AQUA),
                                 ("DeBERTa-v3, fully fine-tuned", "one shared model with two outputs: claim type and arousal", ORANGE),
                                 ("LLM zero-shot (Qwen2.5-1.5B)", "no training at all, just our written category definitions", BLUE)]):
    y = Inches(2.4 + i * 1.3)
    box(s, Inches(0.85), y, Inches(0.22), Inches(0.9), col, radius=False)
    text(s, Inches(1.25), y, Inches(5.2), Inches(0.45), h, size=16, bold=True)
    text(s, Inches(1.25), y + Inches(0.42), Inches(5.2), Inches(0.5), b, size=14, color=MUTED)
box(s, Inches(6.95), Inches(1.6), Inches(5.8), Inches(5.0))
text(s, Inches(7.2), Inches(1.75), Inches(5.3), Inches(0.5), "Arousal from audio: WavLM, partially fine-tuned", size=17, bold=True)
for k in range(12):
    y = Inches(6.05 - k * 0.3)
    frozen = k < 6
    box(s, Inches(7.3), y, Inches(2.0), Inches(0.24), GREY if frozen else RED, radius=False)
    text(s, Inches(9.4), y - Inches(0.05), Inches(0.7), Inches(0.3), f"L{k + 1}", size=10, color=MUTED)
text(s, Inches(10.0), Inches(4.4), Inches(2.6), Inches(1.5), [[("Lower 6 layers frozen", {"bold": True, "color": MUTED})],
     "they keep what it already knows about sound"], size=13)
text(s, Inches(10.0), Inches(2.55), Inches(2.6), Inches(1.6), [[("Upper 6 layers trained", {"bold": True, "color": RED})],
     "retrained on our data, with a learned weight per layer; out comes calm, raised or shouting"], size=13)

s = new_slide(section="Methodology")
title(s, "Checking the car, and testing it fairly")
card(s, Inches(0.6), Inches(1.6), Inches(4.0), Inches(2.4), "Rule-based check",
     "Written before we labelled anything. For example, a tyre claim counts as backed up if laps were at least 0.3 s slower than expected.", size=14)
card(s, Inches(0.6), Inches(4.2), Inches(4.0), Inches(2.4), "Neural check",
     "A 1-D CNN reads the car data, and each claim type looks for its own evidence through gated attention. We compare it with simply joining the inputs.", size=14)
grid = [("Race-grouped splits", "a race is never in both training and testing"), ("Unseen drivers", "we test on drivers the model never saw"),
        ("Held-out season", "train up to 2023, test on 2025 just once"), ("Placebo control", "repeat every check three laps before the message"),
        ("Real chance level", "we simulate random guessing instead of assuming it"), ("Effect size first", "we work out the smallest detectable effect first")]
for i, (h, b) in enumerate(grid):
    col, row = i % 2, i // 2
    x, y = Inches(4.85 + col * 4.0), Inches(1.6 + row * 1.68)
    box(s, x, y, Inches(3.8), Inches(1.5), WHITE, line=LINE)
    text(s, x + Inches(0.18), y + Inches(0.15), Inches(3.4), Inches(0.45), h, size=16, bold=True, color=RED)
    text(s, x + Inches(0.18), y + Inches(0.6), Inches(3.4), Inches(0.85), b, size=14, color=MUTED)

# 05 Problems
section(5, "Problems faced", "What went wrong along the way, and how we fixed each one")

s = new_slide(section="Problems faced")
title(s, "Engineering problems → solutions")
table(s, Inches(0.6), Inches(1.5), Inches(12.1), [
    ["Problem", "What we did", "Outcome"],
    ["Speaker separation would have taken about 19 hours: the model needed 10.9 GB on our 8 GB GPU", "We found the memory overflow and cut the batch size from 32 to 4", "Same output, about 49× faster than real time"],
    ["The official telemetry API, OpenF1, has no data before 2023", "We switched to FastF1, which reads the same feed for every season", f"Data for {ok_races} of 149 races"],
    ["Downloads were capped at 500 requests an hour", "The script waits and retries on its own", "It ran unattended overnight"],
    ["The same “Thank you.” clip appeared in 204 different races", "We removed every copy", "No clip can leak between training and testing"],
    ["Radio timestamps run late", f"We measured it: {pct(lag['A_box_calls']['share_after_pit_entry'])} of “box” calls are stamped after the car is in the pits", "We match messages to laps, not exact seconds"],
    ["We had no access to a paid LLM API", "We ran an open model, Qwen2.5-1.5B, on our own GPU", "That made the zero-shot LLM possible"],
], [4.6, 4.1, 3.4], size=13, row_h=0.66)

s = new_slide(section="Problems faced")
title(s, "Research problems → solutions")
table(s, Inches(0.6), Inches(1.5), Inches(12.1), [
    ["Problem", "What we did"],
    [f"Speaker identification only reached {pct(next(v for k, v in attr['systems'].items() if k.startswith('ensemble_silk (primary'))['accuracy'])}, against a target of 85%",
     "We reported it honestly, and trained every model only on messages a human confirmed were the driver"],
    ["Only 317 verified driver messages, and some claim types were rare", "We kept classes with at least 30 examples, and checked the result still holds at 20"],
    ["The 2025 test season has just 37 items", "Race-grouped cross-validation is our headline; 2025 is reported once as a check"],
    ["Random guessing scores higher than you'd expect on small classes", "We simulated the real chance level for every score"],
    ["Data leakage could make scores look better than they are", "Whole races stay in one fold, and 30 runs on shuffled labels scored exactly chance (0.499)"],
    ["Training on our own rules would make the telemetry model circular", "Instead, the neural model predicts the human label from the car data"],
    ["Broadcasters choose the dramatic moments to air", "A placebo: we repeat every car check three laps before the message"],
], [4.9, 7.2], size=13, row_h=0.6)

# 06 Results
section(6, "Results", "What worked, what didn't, and how sure we are")

s = new_slide(section="Results")
title(s, "Speech: transcription and speaker identification")
ws = wer["systems"]
bar_chart(s, Inches(0.5), Inches(1.6), Inches(5.6), Inches(3.2), ["Dataset transcripts", "Whisper large-v3", "Whisper + F1 prompt"],
          [ws["dataset_cohere"]["wer"], ws["whisper_plain"]["wer"], ws["whisper_prompted"]["wer"]], [GREY, GREY, BLUE], fmt="0%",
          vmax=0.4, title_txt="Word error rate (lower is better)")
bullets(s, Inches(0.6), Inches(5.0), Inches(5.5), Inches(1.8), [
    f"The prompt only cut errors by about {pct(ws['whisper_prompted']['relative_improvement_vs_dataset'])}, well short of our 15% target.",
    "About 3 words in 10 are still wrong on radio audio.",
], size=15)
image(s, "fig1_attribution.png", Inches(6.4), Inches(1.6), w=Inches(6.4))
bullets(s, Inches(6.5), Inches(4.75), Inches(6.2), Inches(2), [
    "Telling driver from engineer: 67% correct, below our 85% target.",
    "The voice helps far more than the words do: 65% vs 39%.",
], size=15)

s = new_slide(section="Results")
title(s, "Main result: text identifies what the driver claims")
mods = [("Majority class", cv["majority"]["claim"]["macro"]["auc_pr"], GREY), ("Context only (no text)", cv["context"]["claim"]["macro"]["auc_pr"], GREY),
        ("Fine-tuned DeBERTa", cv["deberta"]["claim"]["macro"]["auc_pr"], ORANGE), ("TF-IDF", cv["tfidf"]["claim"]["macro"]["auc_pr"], AQUA),
        ("LLM zero-shot", llm["metrics"]["macro"]["auc_pr"], BLUE)]
bar_chart(s, Inches(0.5), Inches(1.5), Inches(6.6), Inches(4.6), [m[0] for m in mods], [m[1] for m in mods], [m[2] for m in mods],
          vmax=0.75, title_txt="Macro AUC-PR, 5 claim types (higher is better)")
text(s, Inches(0.6), Inches(6.15), Inches(6.4), Inches(0.5),
     f"Chance level ≈ {cv['deberta']['claim']['macro']['chance_auc_pr']:.2f}; race-grouped 5-fold cross-validation", size=12, color=MUTED)
p = txt["protocols"]
table(s, Inches(7.4), Inches(1.6), Inches(5.3), [
    ["Split", "DeBERTa", "TF-IDF", "Context"],
    ["Cross-validation", f"{cv['deberta']['claim']['macro']['auc_pr']:.2f}", f"{cv['tfidf']['claim']['macro']['auc_pr']:.2f}", f"{cv['context']['claim']['macro']['auc_pr']:.2f}"],
    ["Test season 2025", f"{p['test2025']['deberta']['claim']['macro']['auc_pr']:.2f}", f"{p['test2025']['tfidf']['claim']['macro']['auc_pr']:.2f}", f"{p['test2025']['context']['claim']['macro']['auc_pr']:.2f}"],
    ["Unseen drivers", f"{p['lodo']['deberta']['claim']['macro']['auc_pr']:.2f}", f"{p['lodo']['tfidf']['claim']['macro']['auc_pr']:.2f}", "–"],
], [2.0, 1.1, 1.1, 1.1], size=14, row_h=0.5)
d = cv["deberta_minus_context"]["claim"]
bullets(s, Inches(7.4), Inches(3.9), Inches(5.3), Inches(2.6), [
    [("Reading the words helps: ", {"bold": True}), (f"+{d['delta_macro_auc_pr']:.2f} over the context-only model (95% CI {d['ci95'][0]:+.2f} to {d['ci95'][1]:+.2f}).", {})],
    "It holds on the unseen 2025 season, and on drivers the model never saw.",
    "With this little data, DeBERTa and simple word counts do about equally well.",
], size=15)

s = new_slide(section="Results")
title(s, "Surprise: a small LLM, told the rules, beats fine-tuning")
image(s, "fig9_learning_curve.png", Inches(0.5), Inches(1.6), w=Inches(6.6))
pv = llm["paired_vs_cv5_oof"]
stat(s, Inches(7.5), Inches(1.6), Inches(2.5), Inches(1.6), f"+{pv['llm_minus_deberta']['delta_macro_auc_pr']:.2f}", "LLM over fine-tuned DeBERTa")
stat(s, Inches(10.2), Inches(1.6), Inches(2.5), Inches(1.6), f"+{pv['llm_minus_tfidf']['delta_macro_auc_pr']:.2f}", "LLM over TF-IDF")
bullets(s, Inches(7.5), Inches(3.5), Inches(5.2), Inches(3.2), [
    "The LLM never saw a single labelled example, only our written category definitions.",
    "With only about 250 examples, telling a model the rule works better than making it learn the rule.",
    "The learning curve on the left: DeBERTa improves fastest as data grows, but flattens well below the LLM.",
    "We wrote the prompt once and never tuned it.",
], size=15)

s = new_slide(section="Results")
title(s, "Where the models go wrong")
image(s, "fig8_confusion.png", Inches(0.3), Inches(1.35), w=Inches(9.0))
bullets(s, Inches(9.0), Inches(1.7), Inches(3.8), Inches(5), [
    "When a driver describes an incident through another driver's behaviour, it gets read as a pace complaint.",
    "Brake complaints, which we can't check, get mistaken for tyre complaints.",
    "Some mistakes start earlier: “Enjoy one more tea guys” is a transcription error, not a model error.",
    "TF-IDF is good at keywords like “contact” and “penalty”; DeBERTa is better on pace.",
], size=15)

s = new_slide(section="Results")
title(s, "Arousal: you can hear it, and fine-tuning helps")
ar_models = [("Majority class", aro["cv5"]["majority"]["macro"]["auc_pr"], GREY),
             ("Audio, frozen WavLM", aro["cv5"]["audio"]["macro"]["auc_pr"], GREY),
             ("Text, TF-IDF", aro["cv5"]["text"]["macro"]["auc_pr"], AQUA),
             ("Audio + text", aro["cv5"]["fused"]["macro"]["auc_pr"], GREY),
             ("Text, DeBERTa multi-task", cv["deberta"]["arousal"]["macro"]["auc_pr"], ORANGE),
             ("Audio, WavLM fine-tuned", ft["protocols"]["cv5"]["metrics"]["macro"]["auc_pr"], RED)]
bar_chart(s, Inches(0.5), Inches(1.5), Inches(6.6), Inches(4.8), [m[0] for m in ar_models], [m[1] for m in ar_models],
          [m[2] for m in ar_models], vmax=0.7, title_txt="Macro AUC-PR, calm / raised / shouting")
lw = ft["full_fit_layer_weights"]
bar_chart(s, Inches(7.5), Inches(1.5), Inches(5.2), Inches(2.7), [k.replace("layer_", "L") for k in lw], list(lw.values()),
          [RED] * len(lw), fmt="0.00", horizontal=False, vmax=0.22, label_size=10, title_txt="Learned weight per WavLM layer")
dd = abl["ft_minus_frozen"]
bullets(s, Inches(7.5), Inches(4.35), Inches(5.2), Inches(2.5), [
    [("Fine-tuning adds ", {"bold": True}), (f"+{dd['delta_macro_auc_pr']:.2f} over frozen WavLM (95% CI {dd['ci95'][0]:+.2f} to {dd['ci95'][1]:+.2f}).", {})],
    f"It still works on drivers it never heard: {ft['protocols']['lodo']['metrics']['macro']['auc_pr']:.2f}, against {aro['cv5']['audio']['macro']['chance_auc_pr']:.2f} for chance.",
    "It leans slightly on the upper-middle layers, 10 and 11.",
], size=15)

s = new_slide(section="Results")
title(s, "Does the car back up the claim? Not beyond chance")
cl = [c for c in ["GRIP_DEG", "TRAFFIC", "INCIDENT", "POWER"] if t2r["claims"][c].get("thresholds")]
prim = {c: t2r["claims"][c]["thresholds"][t2r["claims"][c]["primary_rule"]] for c in cl}
clustered(s, Inches(0.5), Inches(1.5), Inches(7.2), Inches(4.9), [f"{nice[c]} (n={t2r['claims'][c]['n_in_race']})" for c in cl], [
    ("At the claim", [prim[c]["consistency"] for c in cl]), ("Car's usual rate", [prim[c]["base_rate_matched"] for c in cl]),
    ("3 laps before (placebo)", [prim[c]["placebo_rate"] for c in cl])], [RED, GREY, AQUA], vmax=0.8)
gm = t2n["models"]
bullets(s, Inches(8.0), Inches(1.7), Inches(4.8), Inches(5), [
    "How often the car's data “agrees” with a claim, compared with normal laps and with three laps earlier.",
    "No claim type is backed up more often than usual; every confidence interval includes zero.",
    "Incidents look twice as common at a claim, but they were just as common before the driver spoke.",
    f"The neural model agrees: {gm['neural_gated']['macro']['auc_pr']:.2f}, the same as its placebo ({gm['placebo_neural_gated']['macro']['auc_pr']:.2f}) and the context-only model ({gm['context']['macro']['auc_pr']:.2f}).",
    [("What's going on: ", {"bold": True}), ("TV airs radio at eventful moments, and the placebo caught it.", {})],
], size=14, gap=6)

s = new_slide(section="Results")
title(s, "Does a worked-up driver then drive differently?")
image(s, "fig7_t4.png", Inches(0.5), Inches(1.7), w=Inches(7.4))
y1 = t4["outcomes"]["Y1"]
stat(s, Inches(8.3), Inches(1.7), Inches(4.4), Inches(1.6), f"±{t4['mde']['Y1']['mde_s_per_unit_x']:.2f} s", "smallest effect per lap we could detect, computed before fitting")
bullets(s, Inches(8.3), Inches(3.6), Inches(4.4), Inches(3), [
    f"Going from calm to worked up changes lap time by {y1['beta_x_mixed']:+.2f} s (95% CI {y1['ci95_mixed'][0]:+.2f} to {y1['ci95_mixed'][1]:+.2f}).",
    "That isn't significant, and a check with our hand labels shows no effect either.",
    "We also found no sign that some drivers react differently from others.",
    [("In short: ", {"bold": True}), ("a clean null result, reported with its limits.", {})],
], size=15)

s = new_slide(section="Results")
title(s, "Results at a glance")
G, Wk, N = RGBColor(0xE6, 0xF3, 0xEC), RGBColor(0xFF, 0xF4, 0xD6), RGBColor(0xFB, 0xE9, 0xEA)
rows = [["Question", "Result", "Verdict"],
        ["Claim type from text (T1)", f"DeBERTa {cv['deberta']['claim']['macro']['auc_pr']:.2f} vs context {cv['context']['claim']['macro']['auc_pr']:.2f}; holds on unseen drivers and 2025", "Strong"],
        ["Zero-shot LLM vs fine-tuning", f"LLM {llm['metrics']['macro']['auc_pr']:.2f}, +{pv['llm_minus_deberta']['delta_macro_auc_pr']:.2f} over DeBERTa", "Strong"],
        ["Arousal from audio (T3)", f"WavLM fine-tuned {ft['protocols']['cv5']['metrics']['macro']['auc_pr']:.2f}, +{dd['delta_macro_auc_pr']:.2f} over frozen", "Strong"],
        ["Lap-time model", f"{resid['oos_rmse_s_median_race']:.2f} s error vs {resid['naive_driver_mean_rmse_s_median_race']:.2f} s baseline", "Strong"],
        ["ASR prompting", f"Only about a {pct(ws['whisper_prompted']['relative_improvement_vs_dataset'])} gain, against a 15% target", "Weak"],
        ["Speaker identification", "67% correct, against an 85% target", "Failed gate"],
        ["Telemetry agrees with claims (T2)", "No better than normal laps, and the placebo scores the same", "Null"],
        ["Arousal → next laps (T4)", f"{y1['beta_x_mixed']:+.2f} s, smaller than the ±{t4['mde']['Y1']['mde_s_per_unit_x']:.2f} s we could detect", "Null"]]
fills = [None] + [[None, None, G if r[2] == "Strong" else Wk if r[2] == "Weak" else N] for r in rows[1:]]
table(s, Inches(0.6), Inches(1.45), Inches(12.1), rows, [3.8, 6.4, 1.9], size=14, row_h=0.56, fills=fills)

# 07 Limitations & future
section(7, "Limitations and future scope", "Being honest about the edges, and where to go next")

s = new_slide(section="Limitations and future scope")
title(s, "Limitations")
lims = [("One annotator", "One person labelled everything, so we can't measure agreement between labellers.", "checked consistency: 90% on repeats"),
        ("Small data", "Just 317 driver messages, and only 37 in the 2025 test season.", "confidence intervals everywhere"),
        ("Speaker step", "Only 67% accurate, so unlabelled clips still hold some engineer speech.", "trained only on verified driver speech"),
        ("Missing sensors", "Public data has no brake, tyre-temperature or steering readings.", "a separate “can't be checked” class"),
        ("Broadcast selection", "TV producers choose which radio to air, so it isn't a random sample.", "the placebo control exposes it"),
        ("Label influence", "The labeller saw the transcripts, which may favour text-based arousal.", "audio-only model reported separately")]
for i, (h, b_, m) in enumerate(lims):
    col, row = i % 3, i // 3
    x, y = Inches(0.6 + col * 4.1), Inches(1.55 + row * 2.65)
    box(s, x, y, Inches(3.9), Inches(2.45))
    text(s, x + Inches(0.2), y + Inches(0.15), Inches(3.5), Inches(0.45), h, size=18, bold=True)
    text(s, x + Inches(0.2), y + Inches(0.62), Inches(3.5), Inches(1.0), b_, size=15, color=MUTED)
    text(s, x + Inches(0.2), y + Inches(1.65), Inches(3.5), Inches(0.7), [[("What we did: ", {"bold": True, "color": GOOD}), (m, {"color": INK})]], size=14)

s = new_slide(section="Limitations and future scope")
title(s, "Future scope")
fut = [("More and better labels", "A second annotator to measure agreement, and more examples of rare claims like power and traffic."),
       ("LLM-assisted labelling", "Let the LLM suggest labels and have people correct them: faster, and our results suggest it's accurate."),
       ("Better speaker identification", "Train a proper driver-vs-engineer classifier on our hand-labelled turns."),
       ("Fine-tune speech recognition", "Fine-tune Whisper on hand-typed transcripts instead of only prompting it."),
       ("Richer telemetry", "Teams' own data, like tyre temperatures and brake pressure, would make many more claims checkable."),
       ("Other domains", "The same pipeline could work for aviation, emergency-services or industrial radio.")]
for i, (h, b) in enumerate(fut):
    col, row = i % 2, i // 2
    x, y = Inches(0.6 + col * 6.15), Inches(1.55 + row * 1.75)
    box(s, x, y, Inches(5.95), Inches(1.55))
    text(s, x + Inches(0.25), y + Inches(0.15), Inches(0.6), Inches(1.2), f"{i + 1}", size=30, font=HEAD, bold=True, color=RED,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(0.95), y + Inches(0.18), Inches(4.85), Inches(0.45), h, size=17, bold=True)
    text(s, x + Inches(0.95), y + Inches(0.65), Inches(4.85), Inches(0.85), b, size=14, color=MUTED)

# 08 Conclusion
section(8, "Conclusion", "What we learned, in short")
s = new_slide(section="Conclusion")
title(s, "Conclusion")
bullets(s, Inches(0.6), Inches(1.6), Inches(7.2), Inches(5), [
    "We built a full pipeline on real F1 radio, from raw audio all the way to the car's data.",
    [("The words work: ", {"bold": True}), ("you can tell what a driver is claiming from what they say.", {})],
    [("Rules beat a handful of labels: ", {"bold": True}), ("a small LLM given our definitions did better than fine-tuning.", {})],
    [("You can hear arousal, ", {"bold": True}), ("and partially fine-tuning WavLM clearly helps.", {})],
    [("The car rarely backs up a claim ", {"bold": True}), ("once you apply a placebo control.", {})],
    "We report the negative results too; that's what makes the positive ones believable.",
], size=18, gap=12)
box(s, Inches(8.2), Inches(1.6), Inches(4.5), Inches(4.9), DARK)
text(s, Inches(8.5), Inches(1.85), Inches(4), Inches(0.5), "IN ONE SENTENCE", size=13, bold=True, color=RED)
text(s, Inches(8.5), Inches(2.4), Inches(3.9), Inches(3.8),
     "Machines can understand what F1 drivers claim, best when simply told the rules, but public telemetry can't yet verify those claims.",
     size=21, font=HEAD, color=WHITE)

s = new_slide(section="References")
title(s, "References")
refs = [
    "Bredin, H. (2023). pyannote.audio 2.1 speaker diarization pipeline. Interspeech.",
    "Brown, T. et al. (2020). Language models are few-shot learners. NeurIPS.",
    "Chen, S. et al. (2022). WavLM: large-scale self-supervised pre-training for full stack speech processing. IEEE JSTSP 16(6).",
    "Ciletti, M. (2025). F1 Team Radio Dataset. Hugging Face: MikCil/f1-team-radio.",
    "He, P., Gao, J., Chen, W. (2023). DeBERTaV3. ICLR.",
    "Lee, S. et al. (2026). Rethinking multimodal fusion for time series: text modalities need constrained fusion. arXiv:2603.22372.",
    "Nosek, B. et al. (2018). The preregistration revolution. PNAS 115(11).",
    "Qwen Team (2024). Qwen2.5 technical report. arXiv:2412.15115.",
    "Radford, A. et al. (2023). Robust speech recognition via large-scale weak supervision (Whisper). ICML.",
    "Russell, J. A. (1980). A circumplex model of affect. J. Personality and Social Psychology 39(6).",
    "Simmons, J., Nelson, L., Simonsohn, U. (2011). False-positive psychology. Psychological Science 22(11).",
    "Thorne, J. et al. (2018). FEVER: a large-scale dataset for fact extraction and verification. NAACL-HLT.",
    "Wang, H. et al. (2023). WeSpeaker. ICASSP.  ·  Yang, S. et al. (2021). SUPERB. Interspeech.  ·  FastF1 (software).",
]
text(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(5.0), refs, size=13, color=INK)
text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.4), "Full reference list (25 entries) in the IEEE paper; PDFs and links in references/README.md", size=13, italic=True, color=MUTED)

s = new_slide(dark=True)
text(s, Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.2), "Thank you", size=60, font=HEAD, bold=True, color=WHITE)
text(s, Inches(0.8), Inches(3.7), Inches(11.5), Inches(0.8), "Questions?", size=28, color=GREY)
text(s, Inches(0.8), Inches(5.0), Inches(11.5), Inches(0.8),
     "Live demo: one radio message through the whole system (hand-picked examples, for illustration)", size=16, color=GREY)

prs.save(OUT)
print(f"{state['n']} slides -> {OUT}")
